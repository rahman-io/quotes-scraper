# Runs the scraper against a tiny local copy of quotes.toscrape.com
# (html copied from the real site) so the tests don't hit the internet.

import http.server
import json
import subprocess
import sys
import threading
from functools import partial
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))
import scraper

QUOTE = '''<div class="quote"><span class="text">“{text}”</span>
<span>by <small class="author">{author}</small> <a href="/author/{slug}">(about)</a></span>
<div class="tags"> Tags: {tags}</div></div>'''

AUTHOR = '''<html><body><div class="author-details"><h3 class="author-title">{name}</h3>
<p><strong>Born:</strong> <span class="author-born-date">{born}</span> <span class="author-born-location">{place}</span></p>
<div class="author-description"> {bio} </div></div></body></html>'''


def quote(text, author, slug, tags):
    links = ''.join(f'<a class="tag" href="/tag/{t}/page/1/">{t}</a>' for t in tags)
    return QUOTE.format(text=text, author=author, slug=slug, tags=links)


EINSTEIN_1 = quote('The world as we have created it is a process of our thinking.', 'Albert Einstein',
                   'Albert-Einstein', ['change', 'thinking', 'world'])
ROWLING = quote('It is our choices, Harry, that show what we truly are.', 'J.K. Rowling',
                'J-K-Rowling', ['abilities', 'choices'])
EINSTEIN_2 = quote('Try not to become a man of success.', 'Albert Einstein', 'Albert-Einstein', [])
ROWLING_AGAIN = ROWLING  # same quote on page 2, should only be saved once
EINSTEIN_3 = quote('Imagination is more important than knowledge.', 'Albert Einstein',
                   'Albert-Einstein', ['change', 'thinking'])

SITE = {
    'index.html': f'<html><body>{EINSTEIN_1}{ROWLING}'
                  '<ul class="pager"><li class="next"><a href="/page/2/">Next</a></li></ul></body></html>',
    'page/2/index.html': f'<html><body>{EINSTEIN_2}{ROWLING_AGAIN}{EINSTEIN_3}</body></html>',
    'author/Albert-Einstein/index.html': AUTHOR.format(
        name='Albert Einstein', born='March 14, 1879', place='in Ulm, Germany', bio='In 1879, Albert Einstein...'),
    'author/J-K-Rowling/index.html': AUTHOR.format(
        name='J.K. Rowling', born='July 31, 1965',
        place='in Yate, South Gloucestershire, England, The United Kingdom', bio='See also: Robert Galbraith'),
}


@pytest.fixture(scope='module')
def site(tmp_path_factory):
    folder = tmp_path_factory.mktemp('site')
    for path, html in SITE.items():
        (folder / path).parent.mkdir(parents=True, exist_ok=True)
        (folder / path).write_text(html, encoding='utf-8')

    handler = partial(http.server.SimpleHTTPRequestHandler, directory=str(folder))
    handler.log_message = lambda *args: None
    server = http.server.ThreadingHTTPServer(('127.0.0.1', 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f'http://127.0.0.1:{server.server_port}/'
    server.shutdown()


@pytest.fixture(scope='module')
def quotes(site):
    return scraper.scrape_quotes(site, delay=0)


def test_quotes(quotes, site):
    assert len(quotes) == 4  # 5 on the pages, one is a duplicate
    assert quotes[0] == {
        'text': 'The world as we have created it is a process of our thinking.',
        'author': 'Albert Einstein',
        'author_url': site + 'author/Albert-Einstein',
        'tags': ('change', 'thinking', 'world'),
    }
    assert quotes[2]['tags'] == ()


def test_authors(quotes, site):
    authors = scraper.scrape_authors({q['author_url'] for q in quotes}, workers=2)
    assert len(authors) == 2
    assert authors[site + 'author/J-K-Rowling'] == {
        'name': 'J.K. Rowling',
        'born': '1965-07-31',
        'birthplace': 'Yate, South Gloucestershire, England, The United Kingdom',
        'country': 'The United Kingdom',
        'bio': 'See also: Robert Galbraith',
    }


def test_stats(quotes, site):
    authors = scraper.scrape_authors({q['author_url'] for q in quotes})
    stats = scraper.build_stats(quotes, authors)

    assert stats['tags'][:2] == [('change', 2), ('thinking', 2)]
    assert stats['tag_pairs'] == [('change', 'thinking', 2)]  # pairs seen only once are left out
    assert stats['countries'] == [('Germany', 1), ('The United Kingdom', 1)]

    einstein = stats['authors'][0]
    assert einstein[:2] == ('Albert Einstein', 3)
    assert einstein[5] == 'change, thinking, world'


def test_full_run(site, tmp_path):
    out = tmp_path / 'output'
    subprocess.run([sys.executable, str(ROOT / 'scraper.py'), '--url', site, '--delay', '0', '--out', str(out)],
                   check=True, capture_output=True)

    data = json.loads((out / 'quotes.json').read_text(encoding='utf-8'))
    assert len(data['quotes']) == 4 and len(data['authors']) == 2
    assert data['quotes'][1]['author_born'] == '1965-07-31'

    wb = load_workbook(out / 'quotes.xlsx')
    assert wb.sheetnames == ['Quotes', 'Authors', 'Tags', 'Tag pairs', 'Countries']
    assert wb['Quotes'].max_row == 5
    assert wb['Quotes']['C2'].value == 'Germany'
