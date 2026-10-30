import argparse
import json
import os
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from itertools import combinations
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE_URL = 'https://quotes.toscrape.com/'

session = requests.Session()
session.headers['User-Agent'] = (
    'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 '
    '(KHTML, like Gecko) Chrome/130.0 Safari/537.36'
)


def get_soup(url, retries=3):
    for attempt in range(retries):
        try:
            r = session.get(url, timeout=20)  # author pages redirect to http://, requests follows it
            r.raise_for_status()
            r.encoding = 'utf-8'  # otherwise the curly quotes “ ” come out as â€œ
            return BeautifulSoup(r.text, 'html.parser')
        except requests.RequestException as e:
            if attempt == retries - 1:
                raise
            print(f'  {e}, retrying...')
            time.sleep(2)


def parse_quote(div, page_url):
    return {
        'text': div.select_one('span.text').text.strip('“”'),
        'author': div.select_one('small.author').text.strip(),
        'author_url': urljoin(page_url, div.select_one('a[href*="/author/"]')['href']),
        # tuple instead of list: the tags of a quote never change and tuples can go in sets / dict keys
        'tags': tuple(a.text for a in div.select('a.tag')),
    }


def scrape_quotes(start_url, delay=0.5):
    quotes = []
    seen = set()  # (author, text) pairs, so the same quote is never saved twice
    url = start_url
    page = 1

    while url:
        soup = get_soup(url)
        new = 0
        for div in soup.select('div.quote'):
            quote = parse_quote(div, url)
            key = (quote['author'], quote['text'])
            if key not in seen:
                seen.add(key)
                quotes.append(quote)
                new += 1
        print(f'Page {page}: {new} quotes')

        next_page = soup.select_one('li.next a')
        url = urljoin(url, next_page['href']) if next_page else None
        page += 1
        time.sleep(delay)
    return quotes


def parse_author(soup):
    born = soup.select_one('span.author-born-date').text.strip()
    place = soup.select_one('span.author-born-location').text.strip().removeprefix('in ')
    return {
        'name': soup.select_one('h3.author-title').text.strip(),
        'born': datetime.strptime(born, '%B %d, %Y').date().isoformat(),
        'birthplace': place,
        'country': place.split(', ')[-1],
        'bio': soup.select_one('div.author-description').text.strip(),
    }


def scrape_authors(urls, workers=5):
    # 100 quotes but only 50 authors, so every author page is downloaded once
    # and a few at the same time to save some waiting
    def fetch(url):
        return url, parse_author(get_soup(url))

    with ThreadPoolExecutor(max_workers=workers) as pool:
        return dict(pool.map(fetch, sorted(urls)))


def build_stats(quotes, authors):
    by_author = defaultdict(list)
    for q in quotes:
        by_author[q['author_url']].append(q)

    tag_counts = Counter(tag for q in quotes for tag in q['tags'])
    # which tags show up together on the same quote, e.g. ('humor', 'life')
    tag_pairs = Counter(pair for q in quotes for pair in combinations(sorted(q['tags']), 2))
    countries = Counter(a['country'] for a in authors.values())

    author_rows = []
    for url, info in sorted(authors.items(), key=lambda item: -len(by_author[item[0]])):
        tags_used = {tag for q in by_author[url] for tag in q['tags']}
        author_rows.append((info['name'], len(by_author[url]), info['born'], info['birthplace'],
                            info['country'], ', '.join(sorted(tags_used)), url))

    return {
        'authors': author_rows,
        'tags': tag_counts.most_common(),
        'tag_pairs': [(a, b, n) for (a, b), n in tag_pairs.most_common() if n > 1],
        'countries': countries.most_common(),
    }


def save_json(quotes, authors, path):
    data = {
        'quotes': [{**q, 'author_born': authors[q['author_url']]['born']} for q in quotes],
        'authors': authors,
    }
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def save_excel(quotes, authors, stats, path):
    wb = Workbook()
    wb.remove(wb.active)

    quote_rows = [(q['text'], q['author'], authors[q['author_url']]['country'], ', '.join(q['tags']))
                  for q in quotes]
    add_sheet(wb, 'Quotes', ('Quote', 'Author', 'Country', 'Tags'), quote_rows, (90, 22, 22, 45))
    add_sheet(wb, 'Authors', ('Author', 'Quotes', 'Born', 'Birthplace', 'Country', 'Tags used', 'URL'),
              stats['authors'], (22, 8, 12, 45, 20, 60, 50))
    add_sheet(wb, 'Tags', ('Tag', 'Quotes'), stats['tags'], (22, 8))
    add_sheet(wb, 'Tag pairs', ('Tag', 'Tag', 'Quotes together'), stats['tag_pairs'], (20, 20, 16))
    add_sheet(wb, 'Countries', ('Country', 'Authors'), stats['countries'], (25, 8))
    wb.save(path)


def add_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append(row)

    for i, width in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='305496')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def main():
    parser = argparse.ArgumentParser(description='Scrape quotes.toscrape.com to JSON and Excel')
    parser.add_argument('--delay', type=float, default=0.5, help='seconds between page requests (default 0.5)')
    parser.add_argument('--workers', type=int, default=5, help='author pages fetched at the same time (default 5)')
    parser.add_argument('--out', default='output', help='output folder (default: output)')
    parser.add_argument('--url', default=BASE_URL, help='site url, mainly used by the tests')
    args = parser.parse_args()

    quotes = scrape_quotes(args.url, args.delay)

    author_urls = {q['author_url'] for q in quotes}
    print(f'\nGetting {len(author_urls)} author pages...')
    authors = scrape_authors(author_urls, args.workers)

    stats = build_stats(quotes, authors)
    os.makedirs(args.out, exist_ok=True)
    save_json(quotes, authors, os.path.join(args.out, 'quotes.json'))
    save_excel(quotes, authors, stats, os.path.join(args.out, 'quotes.xlsx'))

    print(f'\nDone, {len(quotes)} quotes, {len(authors)} authors and {len(stats["tags"])} tags '
          f'saved to {args.out}/')


if __name__ == '__main__':
    main()
